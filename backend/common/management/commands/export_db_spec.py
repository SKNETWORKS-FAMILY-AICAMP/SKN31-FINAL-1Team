# common/management/commands/export_db_spec.py
"""
현재 Django 모델을 introspect 해서 DB 명세서를 Excel(.xlsx)로 뽑는다.

    python manage.py export_db_spec
    python manage.py export_db_spec --output docs/DB_명세서.xlsx

시트 구성:
  1) 테이블 목록      — 앱 / 모델 / 테이블명 / 한글명 / 설명 / PK / 컬럼 수
  2) 컬럼 상세 명세    — 테이블별 전체 컬럼 (타입 / DB타입 / PK·FK / Null / Unique / 기본값 / Choices)

모델이 바뀌면 다시 실행하면 명세서도 갱신된다.
"""

from django.apps import apps
from django.core.management.base import BaseCommand
from django.db import connection

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# 명세서에 담을 앱 (Django 기본 앱 제외)
TARGET_APPS = ["common", "users", "projects", "meetings", "requirements", "tasks"]

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=10)
CELL_FONT = Font(size=10)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(vertical="top", wrap_text=True)
CENTER = Alignment(vertical="center", horizontal="center")

ON_DELETE_KR = {
    "CASCADE": "CASCADE (같이 삭제)",
    "SET_NULL": "SET NULL",
    "PROTECT": "PROTECT (삭제 차단)",
    "SET_DEFAULT": "SET DEFAULT",
    "DO_NOTHING": "DO NOTHING",
    "RESTRICT": "RESTRICT",
}


def first_docline(model):
    doc = (model.__doc__ or "").strip()
    if not doc or doc.startswith(f"{model.__name__}("):
        return ""
    return doc.splitlines()[0].strip()


def field_rows(model):
    """모델의 concrete 필드 + M2M 를 명세 행(dict)으로."""
    rows = []
    fields = list(model._meta.fields) + list(model._meta.many_to_many)
    for f in fields:
        is_fk = f.is_relation and (f.many_to_one or f.one_to_one)
        is_m2m = f.many_to_many

        # DB 컬럼 타입
        try:
            db_type = f.db_type(connection) or ("M2M(중간테이블)" if is_m2m else "-")
        except Exception:
            db_type = "-"

        # FK 대상
        rel = ""
        on_delete = ""
        if is_fk or is_m2m:
            rt = f.related_model
            rel = f"{rt._meta.db_table}.{rt._meta.pk.column}"
            od = getattr(getattr(f, "remote_field", None), "on_delete", None)
            if od is not None:
                on_delete = ON_DELETE_KR.get(getattr(od, "__name__", ""), getattr(od, "__name__", ""))

        # 기본값
        default = ""
        if f.has_default():
            d = f.default
            default = getattr(d, "__name__", None) or repr(d() if callable(d) else d)
        elif is_fk and f.null:
            default = "NULL"

        # choices
        choices = ""
        if getattr(f, "choices", None):
            choices = ", ".join(f"{k}={v}" for k, v in f.choices)

        rows.append({
            "컬럼명": f.column if hasattr(f, "column") else f.name,
            "한글명(verbose_name)": str(getattr(f, "verbose_name", "")),
            "Django 필드": f.get_internal_type(),
            "DB 타입": db_type,
            "길이": getattr(f, "max_length", "") or "",
            "PK": "Y" if f.primary_key else "",
            "FK →": rel,
            "on_delete": on_delete,
            "Null": "Y" if getattr(f, "null", False) else "",
            "Unique": "Y" if getattr(f, "unique", False) else "",
            "기본값": default,
            "Choices": choices,
        })
    return rows


def style_sheet(ws, headers, widths):
    for c, (h, w) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


class Command(BaseCommand):
    help = "Django 모델 기반 DB 명세서를 Excel(.xlsx)로 생성한다."

    def add_arguments(self, parser):
        parser.add_argument("--output", default="docs/DB_명세서.xlsx", help="저장 경로 (backend 기준 상대경로)")

    def handle(self, *args, **opts):
        wb = Workbook()

        # ── 시트 1: 테이블 목록 ──────────────────────────────
        ws1 = wb.active
        ws1.title = "테이블 목록"
        h1 = ["No.", "앱", "모델", "테이블명", "한글명", "설명", "PK 컬럼", "컬럼 수"]
        w1 = [5, 12, 24, 26, 20, 48, 18, 8]
        style_sheet(ws1, h1, w1)

        models = []
        for app_label in TARGET_APPS:
            try:
                cfg = apps.get_app_config(app_label)
            except LookupError:
                continue
            for m in cfg.get_models():
                models.append((app_label, m))

        for i, (app_label, m) in enumerate(models, start=1):
            meta = m._meta
            ws1.append([
                i, app_label, m.__name__, meta.db_table,
                str(meta.verbose_name), first_docline(m),
                meta.pk.column, len(list(meta.fields)) + len(list(meta.many_to_many)),
            ])
        for row in ws1.iter_rows(min_row=2):
            for cell in row:
                cell.font = CELL_FONT
                cell.alignment = WRAP
                cell.border = BORDER

        # ── 시트 2: 컬럼 상세 명세 ───────────────────────────
        ws2 = wb.create_sheet("컬럼 상세 명세")
        h2 = ["테이블", "No.", "컬럼명", "한글명(verbose_name)", "Django 필드", "DB 타입",
              "길이", "PK", "FK →", "on_delete", "Null", "Unique", "기본값", "Choices"]
        w2 = [24, 5, 22, 24, 16, 20, 6, 5, 26, 18, 5, 7, 18, 40]
        style_sheet(ws2, h2, w2)

        for app_label, m in models:
            table = m._meta.db_table
            for j, r in enumerate(field_rows(m), start=1):
                ws2.append([
                    table, j, r["컬럼명"], r["한글명(verbose_name)"], r["Django 필드"], r["DB 타입"],
                    r["길이"], r["PK"], r["FK →"], r["on_delete"], r["Null"], r["Unique"],
                    r["기본값"], r["Choices"],
                ])
        for row in ws2.iter_rows(min_row=2):
            for cell in row:
                cell.font = CELL_FONT
                cell.alignment = WRAP
                cell.border = BORDER

        # ── 저장 ────────────────────────────────────────────
        from pathlib import Path
        from django.conf import settings
        out = Path(settings.BASE_DIR) / opts["output"]
        out.parent.mkdir(parents=True, exist_ok=True)
        wb.save(out)

        self.stdout.write(self.style.SUCCESS(
            f"DB 명세서 생성 완료: {out}\n"
            f"  - 테이블 {len(models)}개 / 시트 2개 (테이블 목록, 컬럼 상세 명세)"
        ))
